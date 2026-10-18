"""Monta a planilha Excel a partir dos CSVs sintéticos brutos.

A limpeza aqui espelha os passos do Power Query em
01-excel-power-query-tabelas-dinamicas/power-query/, de modo que a planilha
já sai pronta para uso (e para tabelas dinâmicas). Os CSVs brutos continuam
em data/raw/ para o exercício de Power Query.
"""
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parent.parent
BRUTO = RAIZ / "data" / "raw"
SAIDA = RAIZ / "01-excel-power-query-tabelas-dinamicas" / "Analise_Gestao_Ativos_TI.xlsx"

F = "Arial"
base = Font(name=F, size=10)
negrito = Font(name=F, size=10, bold=True)
titulo = Font(name=F, size=14, bold=True)
fonte_cab = Font(name=F, size=10, bold=True, color="FFFFFF")
fundo_cab = PatternFill("solid", fgColor="1F3864")
fonte_entrada = Font(name=F, size=10, color="0000FF")
fundo_entrada = PatternFill("solid", fgColor="FFFF00")
fonte_nota = Font(name=F, size=9, italic=True, color="595959")
FMT_DATA = "dd/mm/yyyy"

# ------------------------------------------------------------------ dados
ativos = pd.read_csv(BRUTO / "ativos.csv", dtype=str, keep_default_na=False, encoding="utf-8")
for c in ativos.columns:
    ativos[c] = ativos[c].str.strip()
ativos["status"] = ativos["status"].str.title()      # "pool - disponível" -> "Pool - Disponível"
ativos = ativos.drop_duplicates(subset="id_ativo")
ativos["departamento"] = ativos["departamento"].replace("", "Não atribuído")
ativos["data_compra"] = pd.to_datetime(ativos["data_compra"], errors="coerce")
ativos["fim_garantia"] = pd.to_datetime(ativos["fim_garantia"], errors="coerce")
ativos["custo_aquisicao_brl"] = ativos["custo_aquisicao_brl"].astype(float)
ativos = ativos.sort_values("id_ativo").reset_index(drop=True)

locais = pd.read_csv(BRUTO / "locais.csv", encoding="utf-8")
movimentacoes = pd.read_csv(BRUTO / "movimentacoes.csv", parse_dates=["data_movimentacao"], encoding="utf-8")
incidentes = pd.read_csv(BRUTO / "incidentes.csv", parse_dates=["data_abertura", "data_resolucao"], encoding="utf-8")

wb = Workbook()


def estilo_cabecalho(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = fonte_cab
        cell.fill = fundo_cab
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def adicionar_tabela(ws, nome, ncols, nlinhas):
    ref = f"A1:{get_column_letter(ncols)}{nlinhas + 1}"
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "A2"


def escrever_df(ws, df, colunas_data=(), larguras=None):
    ws.append(list(df.columns))
    for rec in df.itertuples(index=False):
        linha = []
        for v in rec:
            if isinstance(v, pd.Timestamp):
                v = None if pd.isna(v) else v.to_pydatetime()
            elif (pd.isna(v) if not isinstance(v, str) else False):
                v = None
            linha.append(v)
        ws.append(linha)
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = base
    for ci, col in enumerate(df.columns, start=1):
        if col in colunas_data:
            for r in range(2, len(df) + 2):
                ws.cell(row=r, column=ci).number_format = FMT_DATA
        ws.column_dimensions[get_column_letter(ci)].width = (larguras or {}).get(col, 16)


# ------------------------------------------------------------------ Sobre
ws = wb.active
ws.title = "Sobre"
sobre = [
    ("Análise de Gestão de Ativos de TI - Excel, Power Query e Tabelas Dinâmicas", titulo),
    ("", base),
    ("Todos os dados desta planilha são SINTÉTICOS (simulados). Não vêm de nenhuma organização real e não representam nenhuma organização real.", negrito),
    ("", base),
    ("Abas", negrito),
    ("Indicadores - indicadores principais e tabelas-resumo, todos calculados por fórmulas sobre as abas de dados.", base),
    ("Ativos - 5.000 ativos de informática simulados (físicos + máquinas virtuais). Tabela do Excel: tbl_Ativos.", base),
    ("Locais - 12 unidades simuladas. Tabela do Excel: tbl_Locais.", base),
    ("Movimentações - despachos, devoluções, transferências e retiradas do pool. Tabela do Excel: tbl_Movimentacoes.", base),
    ("Incidentes - incidentes de suporte com gravidade e tempo de resolução. Tabela do Excel: tbl_Incidentes.", base),
    ("", base),
    ("Como usar", negrito),
    ("1. Altere as células amarelas da aba Indicadores (data de referência e limites de ciclo de vida); todos os indicadores se recalculam.", base),
    ("2. Crie Tabelas Dinâmicas a partir das tabelas acima (veja docs/guia-tabelas-dinamicas.md).", base),
    ("3. Para reproduzir a limpeza no Power Query, carregue os CSVs brutos de /data/raw com os scripts M de /power-query.", base),
    ("", base),
    ("Cores das células", negrito),
    ("Texto azul sobre fundo amarelo = entrada que você pode alterar. Texto preto = fórmula ou dado.", base),
    ("", base),
    ("Limpeza aplicada ao arquivo bruto de ativos (mesmos passos do script Power Query)", negrito),
    ("Remover espaços extras, padronizar o status em Title Case, remover 25 linhas duplicadas por id_ativo e substituir departamento em branco por 'Não atribuído'.", base),
]
for i, (texto, fonte) in enumerate(sobre, start=1):
    c = ws.cell(row=i, column=1, value=texto)
    c.font = fonte
ws.column_dimensions["A"].width = 140

# ------------------------------------------------------------------ Indicadores (preenchida depois)
kp = wb.create_sheet("Indicadores")

# ------------------------------------------------------------------ Ativos
wa = wb.create_sheet("Ativos")
cols_a = list(ativos.columns)  # 11 colunas A..K
escrever_df(wa, ativos, colunas_data=("data_compra", "fim_garantia"),
            larguras={"id_ativo": 12, "tipo_ativo": 16, "fabricante": 14, "modelo": 20,
                      "status": 20, "id_local": 12, "departamento": 16})
nA = len(ativos)
ultima = nA + 1
extras = ["idade_anos", "estagio_ciclo_vida", "status_garantia", "nome_local", "regiao"]  # L..P
for i, h in enumerate(extras, start=len(cols_a) + 1):
    wa.cell(row=1, column=i, value=h)
for r in range(2, ultima + 1):
    wa.cell(row=r, column=12, value=f'=IF(E{r}="","",ROUND((Indicadores!$B$4-E{r})/365.25,1))')
    wa.cell(row=r, column=13, value=f'=IF(E{r}="","Virtual",IF(L{r}>=Indicadores!$B$6,"Fim de Vida",IF(L{r}>=Indicadores!$B$5,"Envelhecendo","Atual")))')
    wa.cell(row=r, column=14, value=f'=IF(F{r}="","N/A",IF(F{r}<Indicadores!$B$4,"Vencida","Vigente"))')
    wa.cell(row=r, column=15, value=f'=INDEX(Locais!$B$2:$B$13,MATCH(H{r},Locais!$A$2:$A$13,0))')
    wa.cell(row=r, column=16, value=f'=INDEX(Locais!$C$2:$C$13,MATCH(H{r},Locais!$A$2:$A$13,0))')
    for c in range(12, 17):
        wa.cell(row=r, column=c).font = base
    wa.cell(row=r, column=11).number_format = "#,##0.00"
for c, w_ in zip(range(12, 17), (11, 18, 16, 34, 14)):
    wa.column_dimensions[get_column_letter(c)].width = w_
estilo_cabecalho(wa, 16)
adicionar_tabela(wa, "tbl_Ativos", 16, nA)

# ------------------------------------------------------------------ Locais / Movimentações / Incidentes
wl = wb.create_sheet("Locais")
escrever_df(wl, locais, larguras={"id_local": 12, "nome_local": 36, "regiao": 16})
estilo_cabecalho(wl, 3)
adicionar_tabela(wl, "tbl_Locais", 3, len(locais))

wm = wb.create_sheet("Movimentações")
escrever_df(wm, movimentacoes, colunas_data=("data_movimentacao",),
            larguras={"id_movimentacao": 18, "id_ativo": 12, "tipo_movimentacao": 20,
                      "data_movimentacao": 18, "id_local_origem": 16, "id_local_destino": 16,
                      "departamento_solicitante": 24})
estilo_cabecalho(wm, 7)
adicionar_tabela(wm, "tbl_Movimentacoes", 7, len(movimentacoes))

wi = wb.create_sheet("Incidentes")
escrever_df(wi, incidentes, colunas_data=("data_abertura", "data_resolucao"),
            larguras={"id_incidente": 14, "id_ativo": 12, "id_local": 12, "categoria": 14,
                      "gravidade": 12, "data_abertura": 16, "data_resolucao": 16,
                      "horas_resolucao": 16})
estilo_cabecalho(wi, 8)
adicionar_tabela(wi, "tbl_Incidentes", 8, len(incidentes))
nI = len(incidentes) + 1

# ------------------------------------------------------------------ conteúdo da aba Indicadores
A = lambda col: f"Ativos!${col}$2:${col}${ultima}"          # noqa: E731
I = lambda col: f"Incidentes!${col}$2:${col}${nI}"          # noqa: E731

kp["A1"] = "Gestão de Ativos de TI - Resumo de Indicadores (dados sintéticos)"
kp["A1"].font = titulo
kp["A3"] = "Entradas (altere as células amarelas)"
kp["A3"].font = negrito
entradas = [
    (4, "Data de referência", date(2026, 10, 18), FMT_DATA,
     "Data usada para calcular a idade dos ativos e o status da garantia. Use a data de extração dos dados."),
    (5, "Limite de envelhecimento (anos)", 3, "0",
     "Premissa: ativos com 3 anos ou mais são 'Envelhecendo'. Ciclos de renovação típicos: 3 a 5 anos."),
    (6, "Limite de fim de vida (anos)", 5, "0",
     "Premissa: ativos com 5 anos ou mais estão em 'Fim de Vida' e são candidatos à substituição."),
]
for r, rotulo, val, fmt, nota in entradas:
    kp.cell(row=r, column=1, value=rotulo).font = base
    c = kp.cell(row=r, column=2, value=val)
    c.font, c.fill, c.number_format = fonte_entrada, fundo_entrada, fmt
    kp.cell(row=r, column=3, value=nota).font = fonte_nota

kp["A8"] = "Indicadores principais"
kp["A8"].font = negrito
for ci, h in enumerate(["Indicador", "Valor", "Definição"], start=1):
    kp.cell(row=9, column=ci, value=h)
estilo_cabecalho(kp, 3, row=9)

G, J, L, M, N = A("G"), A("J"), A("L"), A("M"), A("N")
EDATE12 = "EDATE($B$4,-12)"
indicadores = [
    ("Total de ativos", f"=COUNTA({A('A')})", "#,##0", "Todos os registros de ativos após a limpeza (físicos + virtuais)."),
    ("Ativos ativos (exceto Baixados)", f'=COUNTA({A("A")})-COUNTIF({G},"Baixado")', "#,##0", "Total menos os Baixados."),
    ("Em Uso", f'=COUNTIF({G},"Em Uso")', "#,##0", "Atribuídos a um usuário ou departamento."),
    ("Em Estoque", f'=COUNTIF({G},"Em Estoque")', "#,##0", "Disponíveis em estoque, fora do pool."),
    ("Em Manutenção", f'=COUNTIF({G},"Em Manutenção")', "#,##0", "Atualmente em manutenção."),
    ("Baixados", f'=COUNTIF({G},"Baixado")', "#,##0", "Desativados (baixa patrimonial)."),
    ("Ativos do pool (físicos + virtuais)", f'=COUNTIF({J},"S")', "#,##0", "Máquinas gerenciadas no pool compartilhado."),
    ("Utilização do pool", f'=IF(B16=0,0,COUNTIF({G},"Pool - Emprestado")/B16)', "0.0%", "Ativos do pool emprestados / total de ativos do pool."),
    ("Idade média dos ativos físicos ativos (anos)", f'=AVERAGEIFS({L},{G},"<>Baixado")', "0.0", "Máquinas virtuais ficam de fora (sem data de compra)."),
    ("% dos ativos físicos ativos em Fim de Vida", f'=COUNTIFS({M},"Fim de Vida",{G},"<>Baixado")/COUNTIFS({M},"<>Virtual",{G},"<>Baixado")', "0.0%", "Idade >= limite de fim de vida."),
    ("% dos ativos físicos ativos fora da garantia", f'=COUNTIFS({N},"Vencida",{G},"<>Baixado")/COUNTIFS({N},"<>N/A",{G},"<>Baixado")', "0.0%", "Data de fim da garantia anterior à data de referência."),
    ("Incidentes nos últimos 12 meses", f'=COUNTIFS({I("F")},">"&{EDATE12},{I("F")},"<="&$B$4)', "#,##0", "Abertos nos 12 meses anteriores à data de referência."),
    ("Incidentes por 100 ativos ativos (12 meses)", "=IF(B11=0,0,B21/B11*100)", "0.0", "Taxa de incidentes normalizada pelo tamanho do parque."),
    ("Tempo médio de resolução (horas)", f'=AVERAGE({I("H")})', "0.0", "Média de horas_resolucao em todos os incidentes."),
]
for i, (rotulo, formula, fmt, definicao) in enumerate(indicadores, start=10):
    kp.cell(row=i, column=1, value=rotulo).font = base
    c = kp.cell(row=i, column=2, value=formula)
    c.font, c.number_format = base, fmt
    kp.cell(row=i, column=3, value=definicao).font = fonte_nota

# Por local
r0 = 26
kp.cell(row=r0 - 1, column=1, value="Por local").font = negrito
cab_local = ["Local", "Ativos ativos", "Idade média (anos)", "Fim de Vida", "Ativos do pool",
             "Incidentes (12m)", "Incidentes por 100 ativos"]
for ci, h in enumerate(cab_local, start=1):
    kp.cell(row=r0, column=ci, value=h)
estilo_cabecalho(kp, 7, row=r0)
for k in range(12):
    r = r0 + 1 + k
    lr = k + 2
    kp.cell(row=r, column=1, value=f"=Locais!B{lr}")
    kp.cell(row=r, column=2, value=f'=COUNTIFS({A("H")},Locais!$A{lr},{G},"<>Baixado")')
    kp.cell(row=r, column=3, value=f'=IFERROR(AVERAGEIFS({L},{A("H")},Locais!$A{lr},{G},"<>Baixado"),0)')
    kp.cell(row=r, column=4, value=f'=COUNTIFS({A("H")},Locais!$A{lr},{M},"Fim de Vida",{G},"<>Baixado")')
    kp.cell(row=r, column=5, value=f'=COUNTIFS({A("H")},Locais!$A{lr},{J},"S")')
    kp.cell(row=r, column=6, value=f'=COUNTIFS({I("C")},Locais!$A{lr},{I("F")},">"&{EDATE12},{I("F")},"<="&$B$4)')
    kp.cell(row=r, column=7, value=f"=IF(B{r}=0,0,F{r}/B{r}*100)")
rt = r0 + 13
kp.cell(row=rt, column=1, value="Total").font = negrito
for col in "BDEF":
    kp[f"{col}{rt}"] = f"=SUM({col}{r0+1}:{col}{rt-1})"
kp[f"C{rt}"] = "=B18"
kp[f"G{rt}"] = f"=IF(B{rt}=0,0,F{rt}/B{rt}*100)"
for r in range(r0 + 1, rt + 1):
    for ci, fmt in zip(range(1, 8), (None, "#,##0", "0.0", "#,##0", "#,##0", "#,##0", "0.0")):
        c = kp.cell(row=r, column=ci)
        c.font = negrito if r == rt else base
        if fmt:
            c.number_format = fmt

# Por tipo de ativo
t0 = rt + 3
kp.cell(row=t0 - 1, column=1, value="Por tipo de ativo (ativos ativos)").font = negrito
for ci, h in enumerate(["Tipo de ativo", "Ativos ativos", "Idade média (anos)", "Fim de Vida", "% Fim de Vida"], start=1):
    kp.cell(row=t0, column=ci, value=h)
estilo_cabecalho(kp, 5, row=t0)
tipos = ["Desktop", "Notebook", "Workstation", "Thin Client", "Máquina Virtual"]
B_ = A("B")
for k, t in enumerate(tipos):
    r = t0 + 1 + k
    kp.cell(row=r, column=1, value=t).font = base
    kp.cell(row=r, column=2, value=f'=COUNTIFS({B_},$A{r},{G},"<>Baixado")')
    kp.cell(row=r, column=3, value=f'=IFERROR(AVERAGEIFS({L},{B_},$A{r},{G},"<>Baixado"),0)')
    kp.cell(row=r, column=4, value=f'=COUNTIFS({B_},$A{r},{M},"Fim de Vida",{G},"<>Baixado")')
    kp.cell(row=r, column=5, value=f"=IF(B{r}=0,0,D{r}/B{r})")
    for ci, fmt in zip(range(2, 6), ("#,##0", "0.0", "#,##0", "0.0%")):
        c = kp.cell(row=r, column=ci)
        c.font, c.number_format = base, fmt

# Por estágio do ciclo de vida
s0 = t0 + len(tipos) + 3
kp.cell(row=s0 - 1, column=1, value="Por estágio do ciclo de vida (ativos ativos)").font = negrito
for ci, h in enumerate(["Estágio do ciclo de vida", "Ativos ativos", "% dos ativos"], start=1):
    kp.cell(row=s0, column=ci, value=h)
estilo_cabecalho(kp, 3, row=s0)
for k, st in enumerate(["Atual", "Envelhecendo", "Fim de Vida", "Virtual"]):
    r = s0 + 1 + k
    kp.cell(row=r, column=1, value=st).font = base
    kp.cell(row=r, column=2, value=f'=COUNTIFS({M},$A{r},{G},"<>Baixado")')
    kp.cell(row=r, column=3, value=f"=IF($B$11=0,0,B{r}/$B$11)")
    kp.cell(row=r, column=2).number_format = "#,##0"
    kp.cell(row=r, column=3).number_format = "0.0%"
    for ci in (2, 3):
        kp.cell(row=r, column=ci).font = base

kp.column_dimensions["A"].width = 48
kp.column_dimensions["B"].width = 16
kp.column_dimensions["C"].width = 18
for col in "DEFG":
    kp.column_dimensions[col].width = 20
kp.freeze_panes = "A3"

wb.save(SAIDA)
print("Salvo em", SAIDA, "| ativos após a limpeza:", nA)
